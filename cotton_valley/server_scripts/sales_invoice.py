import os
import mimetypes
import frappe
from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage

from PIL import Image as PILImage  # ✅ required for webp conversion

mimetypes.add_type("image/webp", ".webp")


@frappe.whitelist()
def download_sales_invoice_excel(sales_invoice):
    doc = frappe.get_doc("Sales Invoice", sales_invoice)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Invoice"

    bold = Font(bold=True)
    title_font = Font(bold=True, size=14)
    head_font = Font(bold=True, size=12)

    left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="top", wrap_text=True)

    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def put(r, c, v, f=None, a=None, b=None):
        cell = ws.cell(row=r, column=c, value=v)
        if f:
            cell.font = f
        if a:
            cell.alignment = a
        if b:
            cell.border = b
        return cell

    def resolve_image_path(image_url: str):
        """Return local path for /files/ or /private/files/ urls. Otherwise None."""
        if not image_url:
            return None
        if image_url.startswith("http://") or image_url.startswith("https://"):
            return None  # skip external
        if image_url.startswith("/private/files/"):
            return frappe.get_site_path(image_url.lstrip("/"))
        if image_url.startswith("/files/"):
            return frappe.get_site_path("public", image_url.lstrip("/"))
        if image_url.startswith("files/"):
            return frappe.get_site_path("public", image_url)
        if image_url.startswith("private/files/"):
            return frappe.get_site_path(image_url)
        return None

    def convert_webp_to_png(path: str):
        """Convert .webp to .png in /tmp and return png path; otherwise return original."""
        if not path:
            return None
        ext = os.path.splitext(path)[1].lower()
        if ext != ".webp":
            return path

        base = os.path.splitext(os.path.basename(path))[0]
        png_path = f"/tmp/{base}.png"

        im = PILImage.open(path).convert("RGBA")
        im.save(png_path, "PNG")
        return png_path

    # =========================
    # HEADER
    # =========================
    r = 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    put(r, 1, (doc.company or "Company"), title_font, center)
    r += 1

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    put(r, 1, doc.doctype, head_font, center)
    r += 2

    put(r, 1, "Billing Information:", bold, left)
    put(r, 4, "Shipping Information:", bold, left)
    put(r, 8, "Order #", bold, left)
    put(r, 9, doc.name, bold, left)
    r += 1

    billing_lines = []
    if doc.get("customer_address"):
        a = frappe.get_doc("Address", doc.customer_address)
        billing_lines = [
            a.get("address_title") or "",
            a.get("address_line1") or "",
            a.get("address_line2") or "",
            " ".join([x for x in [a.get("city"), a.get("state"), a.get("pincode")] if x]),
            a.get("country") or "",
            f"Phone: {a.get('phone')}" if a.get("phone") else "",
        ]
        mobile = a.get("mobile_no") or a.get("mobile") or a.get("mobile_number")
        if mobile:
            billing_lines.append(f"Mobile: {mobile}")

    shipping_lines = []
    # Sales Invoice me shipping address ka field aksar `shipping_address_name`
    if doc.get("shipping_address_name"):
        s = frappe.get_doc("Address", doc.shipping_address_name)
        shipping_lines = [
            s.get("address_title") or "",
            s.get("address_line1") or "",
            s.get("address_line2") or "",
            " ".join([x for x in [s.get("city"), s.get("state"), s.get("pincode")] if x]),
            s.get("country") or "",
            f"Phone: {s.get('phone')}" if s.get("phone") else "",
        ]
        mobile = s.get("mobile_no") or s.get("mobile") or s.get("mobile_number")
        if mobile:
            shipping_lines.append(f"Mobile: {mobile}")

    # ✅ Sales Rep (Invoice me sales_team ho bhi sakta hai, nahi bhi)
    sales_rep_name = ""
    if doc.get("sales_team"):
        try:
            sales_rep_name = doc.sales_team[0].sales_person
        except Exception:
            sales_rep_name = ""

    sales_rep_email = ""
    sales_rep_phone = ""
    if sales_rep_name:
        emp_name = frappe.db.get_value("Employee", {"employee_name": sales_rep_name}, "name")
        if emp_name:
            emp = frappe.get_doc("Employee", emp_name)
            sales_rep_email = emp.get("user_id") or emp.get("company_email") or emp.get("personal_email") or ""
            sales_rep_phone = emp.get("mobile") or emp.get("cell_number") or emp.get("phone_number") or ""

    # ✅ Processed By: show sales rep name (same as your Sales Order requirement)
    processed_by = sales_rep_name or "—"

    # Account # (same fallback logic)
    account_no = (
        doc.get("customer_account_number")
        or doc.get("account_no")
        or doc.get("customer_account")
        or "—"
    )

    # Order Date
    order_date = str(doc.posting_date or "")

    meta_rows = [
        ("Order Date", order_date),
        ("Sales Rep", sales_rep_name or "—"),
        ("Sales Rep Phone", sales_rep_phone or "—"),
        ("Sales Rep Email", sales_rep_email or "—"),
        ("Processed By", processed_by or "—"),
        ("Account #", account_no or "—"),
        ("Status", doc.status or "—"),
        ("Order Type", doc.get("order_type") or doc.get("custom_order_type") or "—"),
    ]

    max_lines = max(len(billing_lines), len(shipping_lines), len(meta_rows))
    for i in range(max_lines):
        btxt = billing_lines[i] if i < len(billing_lines) else ""
        ws.merge_cells(start_row=r + i, start_column=1, end_row=r + i, end_column=3)
        put(r + i, 1, btxt, None, left)

        stxt = shipping_lines[i] if i < len(shipping_lines) else ""
        ws.merge_cells(start_row=r + i, start_column=4, end_row=r + i, end_column=7)
        put(r + i, 4, stxt, None, left)

        if i < len(meta_rows):
            put(r + i, 8, meta_rows[i][0], bold, left)
            put(r + i, 9, meta_rows[i][1], None, left)
        else:
            put(r + i, 8, "", None, left)
            put(r + i, 9, "", None, left)

    r += max_lines + 1

    # ✅ Customer email from Customer doctype field: custom_customer_email
    customer_email = ""
    if doc.get("customer"):
        # try doc field (if exists), otherwise from Customer doctype
        customer_email = doc.get("custom_customer_email") or ""
        if not customer_email:
            customer_email = frappe.db.get_value("Customer", doc.customer, "custom_customer_email") or ""

    put(r, 1, "Email Address:", bold, left)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    put(r, 2, customer_email or "—", None, left)
    r += 2

    # =========================
    # ITEMS TABLE (Image + custom_case_pack)
    # =========================
    headers = ["Line#", "SKU", "Image", "Item Name", "Case Pack", "Qty", "UOM", "Rate", "Amount"]
    for i, h in enumerate(headers, start=1):
        put(r, i, h, bold, center, border)
    r += 1

    line_no = 1
    for it in doc.items:
        case_pack = it.get("custom_case_pack") or ""

        ws.row_dimensions[r].height = 55  # for image

        row_values = [
            line_no,
            it.item_code or "",
            "",  # image inserted
            it.item_name or "",
            case_pack,
            float(it.qty or 0),
            it.uom or "",
            float(it.rate or 0),
            float(it.amount or 0),
        ]

        for c, v in enumerate(row_values, start=1):
            align = left if c in (2, 4) else center
            if c in (8, 9):
                align = right
            put(r, c, v, None, align, border)

        img_path = resolve_image_path(it.get("image"))
        if img_path and os.path.exists(img_path):
            try:
                img_path = convert_webp_to_png(img_path)
                xl_img = XLImage(img_path)
                xl_img.width = 45
                xl_img.height = 45
                ws.add_image(xl_img, f"C{r}")
            except Exception:
                pass

        r += 1
        line_no += 1

    r += 1

    # =========================
    # TOTALS
    # =========================
    put(r, 8, "Subtotal", bold, right)
    put(r, 9, float(doc.total or 0), None, right)
    r += 1

    put(r, 8, "Tax", bold, right)
    put(r, 9, float(doc.total_taxes_and_charges or 0), None, right)
    r += 1

    put(r, 8, "Grand Total", bold, right)
    put(r, 9, float(doc.grand_total or 0), bold, right)

    # Column widths (includes image col)
    widths = [8, 18, 12, 40, 12, 10, 10, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    bio = BytesIO()
    wb.save(bio)

    frappe.response["filename"] = f"{doc.name}.xlsx"
    frappe.response["filecontent"] = bio.getvalue()
    frappe.response["type"] = "binary"
